"""Export CHALLENGE1_FINAL_REPORT.md to DOCX and PDF."""

from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, Preformatted

ROOT = Path(__file__).parent
SOURCE = ROOT / "CHALLENGE1_FINAL_REPORT.md"
DOCX_OUTPUT = ROOT / "results" / "CHALLENGE1_FINAL_REPORT.docx"
PDF_OUTPUT = ROOT / "results" / "CHALLENGE1_FINAL_REPORT.pdf"


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*", line))


def table_cells(line: str):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text


def parse_blocks(lines):
    blocks = []
    index = 0
    while index < len(lines):
        line = lines[index].rstrip("\n")
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            code = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code.append(lines[index].rstrip("\n"))
                index += 1
            blocks.append(("code", "\n".join(code)))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and lines[index + 1].startswith("|") and is_separator(lines[index + 1]):
            rows = [table_cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(table_cells(lines[index]))
                index += 1
            blocks.append(("table", rows))
            continue
        image_match = re.search(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            blocks.append(("image", image_match.group(2)))
            index += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append((f"h{min(level, 3)}", clean_inline(line[level:].strip())))
        elif line.startswith("- "):
            blocks.append(("bullet", clean_inline(line[2:])))
        else:
            blocks.append(("paragraph", clean_inline(line)))
        index += 1
    return blocks


def export_docx(blocks):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)

    for kind, value in blocks:
        if kind.startswith("h"):
            level = int(kind[1:])
            document.add_heading(value, level=level)
        elif kind == "paragraph":
            document.add_paragraph(value)
        elif kind == "bullet":
            document.add_paragraph(value, style="List Bullet")
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(value)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        elif kind == "table":
            table = document.add_table(rows=0, cols=len(value[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(value):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(7 if row_index else 7.5)
        elif kind == "image":
            image_path = ROOT / value
            if image_path.exists():
                document.add_picture(str(image_path), width=Inches(6.4))
    document.save(DOCX_OUTPUT)


def export_pdf(blocks):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportH1", parent=styles["Heading1"], fontSize=15, leading=18, spaceAfter=8))
    styles.add(ParagraphStyle(name="ReportH2", parent=styles["Heading2"], fontSize=11, leading=14, spaceBefore=7, spaceAfter=4))
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontSize=8.2, leading=10.2, spaceAfter=4))
    styles.add(ParagraphStyle(name="ReportCode", parent=styles["Code"], fontSize=6.5, leading=8))
    story = []

    for kind, value in blocks:
        if kind == "h1":
            story.append(Paragraph(value, styles["ReportH1"]))
        elif kind in ("h2", "h3"):
            story.append(Paragraph(value, styles["ReportH2"]))
        elif kind == "paragraph":
            story.append(Paragraph(value.replace("&", "&amp;"), styles["ReportBody"]))
        elif kind == "bullet":
            story.append(Paragraph("&#8226; " + value.replace("&", "&amp;"), styles["ReportBody"]))
        elif kind == "code":
            story.append(Preformatted(value, styles["ReportCode"]))
            story.append(Spacer(1, 3))
        elif kind == "table":
            data = [[Paragraph(cell.replace("&", "&amp;"), styles["ReportBody"]) for cell in row] for row in value]
            table = Table(data, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9eaf0")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(table)
            story.append(Spacer(1, 5))
        elif kind == "image":
            image_path = ROOT / value
            if image_path.exists():
                image = Image(str(image_path))
                image._restrictSize(6.6 * inch, 4.2 * inch)
                story.append(image)
                story.append(Spacer(1, 6))

    document = SimpleDocTemplate(str(PDF_OUTPUT), pagesize=letter, rightMargin=0.55 * inch, leftMargin=0.55 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch)
    document.build(story)


if __name__ == "__main__":
    report_blocks = parse_blocks(SOURCE.read_text(encoding="utf-8").splitlines())
    DOCX_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    export_docx(report_blocks)
    export_pdf(report_blocks)
    print(DOCX_OUTPUT)
    print(PDF_OUTPUT)
